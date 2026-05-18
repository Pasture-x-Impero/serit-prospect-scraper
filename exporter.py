"""
Eksport av resultater til Excel
===============================
Lager en ryddig Excel-fil med alle selskapsdata.
"""

import os
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR, FYLKER

logger = logging.getLogger(__name__)

# Kolonnedefinisjoner
KOLONNER = [
    ("Organisasjonsnummer", 18),
    ("Selskapsnavn", 40),
    ("Org.form", 10),
    ("NACE-kode", 12),
    ("Næringsbeskrivelse", 35),
    ("Antall ansatte", 14),
    ("Stiftelsesdato", 14),
    ("Adresse", 30),
    ("Postnummer", 12),
    ("Poststed", 18),
    ("Fylke", 18),
    ("Nettside", 35),
    ("Nettside kilde", 13),
    ("Daglig leder", 25),
    ("E-post", 30),
    ("Telefon", 16),
    ("Omsetning (MNOK)", 16),
    ("Driftsresultat (MNOK)", 20),
    ("Egenkapital (MNOK)", 18),
    ("Regnskapsår", 13),
]


def eksporter_til_excel(
    data: list[dict],
    fylkesnummer: str,
    statistikk: dict,
    ekskluderte: list[dict] | None = None,
) -> str:
    """
    Eksporter resultater til Excel-fil.

    Args:
        data: Liste med berikede selskapsdata
        fylkesnummer: Fylkesnummer for filnavn
        statistikk: Filtreringsstatistikk

    Returns:
        Filsti til generert Excel-fil
    """
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    fylke_navn = FYLKER.get(fylkesnummer, fylkesnummer).replace(" ", "_")
    tidsstempel = datetime.now().strftime("%Y%m%d_%H%M")
    filnavn = f"serit_kandidater_{fylke_navn}_{tidsstempel}.xlsx"
    filsti = os.path.join(OUTPUT_DIR, filnavn)

    wb = Workbook()

    # --- Ark 1: Kandidater ---
    ws = wb.active
    ws.title = "Kandidater"

    # Stiler
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    alt_fill = PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid")

    # Overskriftsrad
    for col_idx, (header, width) in enumerate(KOLONNER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Frys første rad
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(KOLONNER))}1"

    data = sorted(data, key=lambda x: (x.get("fylke") or "", -(x.get("antall_ansatte") or 0)))

    # Datarad
    felt_mapping = [
        "organisasjonsnummer", "navn", "organisasjonsform",
        "naeringskode", "naeringsbeskrivelse", "antall_ansatte",
        "stiftelsesdato", "adresse", "postnummer", "poststed",
        "fylke", "nettside", "nettside_kilde", "daglig_leder", "epost",
        "telefon",
        "omsetning", "driftsresultat", "egenkapital", "regnskap_aar",
    ]

    nettside_col = felt_mapping.index("nettside") + 1
    ddg_fill = PatternFill(start_color="FFF8E7", end_color="FFF8E7", fill_type="solid")
    ddg_fill_alt = PatternFill(start_color="FFF3D0", end_color="FFF3D0", fill_type="solid")
    hyperlink_font = Font(name="Calibri", size=10, color="0563C1", underline="single")

    for row_idx, selskap in enumerate(data, 2):
        er_ddg = selskap.get("nettside_kilde") == "DDG"
        for col_idx, felt in enumerate(felt_mapping, 1):
            verdi = selskap.get(felt, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=verdi or "")
            cell.border = border
            if col_idx == nettside_col and verdi:
                cell.hyperlink = verdi
                cell.font = hyperlink_font
            else:
                cell.font = cell_font
            if er_ddg:
                cell.fill = ddg_fill_alt if row_idx % 2 == 0 else ddg_fill
            elif row_idx % 2 == 0:
                cell.fill = alt_fill

    # --- Ark 2: Statistikk ---
    ws_stat = wb.create_sheet("Statistikk")
    ws_stat.column_dimensions["A"].width = 35
    ws_stat.column_dimensions["B"].width = 15

    stat_header_font = Font(name="Calibri", bold=True, size=12)
    stat_font = Font(name="Calibri", size=11)

    ws_stat.cell(row=1, column=1, value="Filtreringsrapport").font = Font(
        name="Calibri", bold=True, size=14
    )
    ws_stat.cell(row=2, column=1, value=f"Fylke: {fylke_navn.replace('_', ' ')}").font = stat_font
    ws_stat.cell(row=3, column=1, value=f"Dato: {datetime.now().strftime('%d.%m.%Y %H:%M')}").font = stat_font

    stat_rader = [
        ("", ""),
        ("Filtreringssteg", "Antall"),
        ("Totalt hentet fra Enhetsregisteret¹", statistikk.get("totalt_inn", 0)),
        ("Fjernet: Duplikater", statistikk.get("fjernet_duplikat", 0)),
        ("Fjernet: Serit-selskap", statistikk.get("fjernet_serit", 0)),
        ("Fjernet: Konkurrenter", statistikk.get("fjernet_konkurrent", 0)),
        ("Fjernet: Feil primær NACE-kode", statistikk.get("fjernet_feil_nace", 0)),
        ("Fjernet: For mange ansatte (> 50)", statistikk.get("fjernet_for_mange_ansatte", 0)),
        ("Fjernet: For høy omsetning (> 100 MNOK)", statistikk.get("fjernet_for_stor_omsetning", 0)),
        ("Fjernet: Irrelevante nøkkelord", statistikk.get("fjernet_irrelevant", 0)),
        ("", ""),
        ("Kandidater i resultatet", statistikk.get("totalt_ut", 0)),
        ("", ""),
        ("¹ Forhåndsfiltrert av Brreg API: kun AS/ASA,", ""),
        ("  min. 5 ansatte, aktive selskaper", ""),
    ]

    for i, (label, verdi) in enumerate(stat_rader, 5):
        c1 = ws_stat.cell(row=i, column=1, value=label)
        c2 = ws_stat.cell(row=i, column=2, value=verdi)
        if label in ("Filtreringssteg", "Kandidater i resultatet"):
            c1.font = stat_header_font
            c2.font = stat_header_font
        else:
            c1.font = stat_font
            c2.font = stat_font

    # --- Ark 3: Ekskluderte ---
    if ekskluderte:
        ws_ekskl = wb.create_sheet("Ekskluderte")
        ekskl_kolonner = [
            ("Organisasjonsnummer", 18),
            ("Selskapsnavn", 40),
            ("Antall ansatte", 14),
            ("Grunn", 22),
            ("Detalj", 50),
        ]
        ekskl_felt = ["organisasjonsnummer", "navn", "antall_ansatte", "grunn", "detalj"]

        for col_idx, (header, width) in enumerate(ekskl_kolonner, 1):
            cell = ws_ekskl.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
            ws_ekskl.column_dimensions[get_column_letter(col_idx)].width = width

        ws_ekskl.freeze_panes = "A2"
        ws_ekskl.auto_filter.ref = f"A1:{get_column_letter(len(ekskl_kolonner))}1"

        ekskluderte = sorted(ekskluderte, key=lambda x: x.get("antall_ansatte") or 0, reverse=True)
        for row_idx, selskap in enumerate(ekskluderte, 2):
            for col_idx, felt in enumerate(ekskl_felt, 1):
                verdi = selskap.get(felt, "")
                cell = ws_ekskl.cell(row=row_idx, column=col_idx, value=verdi or "")
                cell.font = cell_font
                cell.border = border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

    wb.save(filsti)
    logger.info(f"Excel-fil lagret: {filsti}")
    return filsti
